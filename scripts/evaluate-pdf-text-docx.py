"""Independently inspect actual editable DOCX output and edit/reopen a copy."""

import argparse
import hashlib
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
import pypdf
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run", type=Path)
    parser.add_argument("--attempt", default="r1")
    args = parser.parse_args()
    criterion = ROOT / "fixtures/pdf-text-docx/criteria.json"
    assert (
        hashlib.sha256(criterion.read_bytes()).hexdigest()
        == "52a3dbd5e25b943939066defc94e51840d7601cb392b75c9bd5df63115d570f9"
    )
    fixed = json.loads(criterion.read_text("utf-8"))
    rows = []
    for name in ("word-text.docx", "word-text-edited.docx", "word-ui-edited.docx"):
        row = dict(file=name, status="FAIL")
        try:
            path = args.run / name
            expected = []
            pages = (
                fixed["expected_pages"][:1]
                if name == "word-ui-edited.docx"
                else fixed["expected_pages"]
            )
            for page, lines in enumerate(pages):
                if page:
                    expected.append("")  # Paragraph containing the explicit page break.
                expected.extend(
                    line.replace("00123", "00456") if "edited" in name else line
                    for line in lines
                )
            with zipfile.ZipFile(path) as archive:
                assert archive.testzip() is None
                assert set(archive.namelist()) == {
                    "[Content_Types].xml",
                    "_rels/.rels",
                    "word/document.xml",
                }
                tree = ET.fromstring(archive.read("word/document.xml"))
                breaks = tree.findall('.//w:br[@w:type="page"]', NS)
                assert len(breaks) == len(pages) - 1
                assert not tree.findall(".//w:instrText", NS)
                assert not tree.findall(".//w:drawing", NS)
                relationships = ET.fromstring(archive.read("_rels/.rels"))
                assert all(
                    item.get("TargetMode") != "External" for item in relationships
                )
            document = Document(path)
            actual = [paragraph.text for paragraph in document.paragraphs]
            assert (
                actual == expected
            ), "Exact editable paragraphs, including page boundaries"
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            document.paragraphs[0].text = "独立読取り後の編集。"
            edited = args.run / (path.stem + "-python-edit-" + args.attempt + ".docx")
            assert not edited.exists()
            document.save(edited)
            reread = Document(edited)
            assert reread.paragraphs[0].text == "独立読取り後の編集。"
            assert [p.text for p in reread.paragraphs[1:]] == expected[1:]
            assert hashlib.sha256(path.read_bytes()).hexdigest() == before
            row.update(
                status="PASS",
                paragraphs=actual,
                explicit_page_breaks=len(breaks),
                sha256=before,
                independent_edit_reopen=True,
                edited_copy_sha256=hashlib.sha256(edited.read_bytes()).hexdigest(),
            )
        except Exception as error:
            row["error"] = str(error)
        rows.append(row)
    row = dict(file="word-engine.pdf", status="FAIL")
    try:
        path = args.run / row["file"]
        reader = pypdf.PdfReader(path)
        observed = []
        with pdfium.PdfDocument(path) as pdf:
            assert len(reader.pages) == len(pdf) == len(fixed["expected_pages"])
            for index, phrases in enumerate(fixed["expected_pages"]):
                page = pdf[index]
                text = page.get_textpage()
                extracted = text.get_text_range()
                other = reader.pages[index].extract_text()
                for phrase in phrases:
                    assert phrase in extracted, (
                        "Exact independent PDFium text: " + phrase
                    )
                    assert phrase in other, "Exact independent pypdf text: " + phrase
                    search = text.search(phrase)
                    match = search.get_next()
                    search.close()
                    assert match, "Independent PDFium search: " + phrase
                    start, count = match
                    assert count > 0 and text.get_charbox(start)[2] > 0
                page.render(scale=1.5).to_pil().save(
                    args.run / f"word-engine-page-{index + 1}-{args.attempt}.png"
                )
                observed.append(extracted)
                text.close()
                page.close()
        row.update(
            status="PASS",
            sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            exact_PDFium_and_pypdf_text=True,
            all_fixed_phrases_searchable=True,
            pages_rendered=len(observed),
            observed=observed,
        )
    except Exception as error:
        row["error"] = str(error)
    rows.append(row)
    result = dict(
        criteria_sha256=hashlib.sha256(criterion.read_bytes()).hexdigest(),
        cases=rows,
        failures=sum(row["status"] != "PASS" for row in rows),
        Word_native_GUI="未実行",
        original_page_layout_restoration="対象外（本文取り出し）",
    )
    with (args.run / ("word-independent-" + args.attempt + ".json")).open(
        "x", encoding="utf-8"
    ) as output:
        json.dump(result, output, ensure_ascii=False, indent=2)
        output.write("\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(bool(result["failures"]))


if __name__ == "__main__":
    main()
