"""Inspect XLSX cells independently of the PDF extraction implementation."""

import argparse
import hashlib
import json
from pathlib import Path
import zipfile

import openpyxl
import pypdfium2 as pdfium
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
CRITERIA_SHA = "07c15ed09f71d8b5f1b5eeff45b65a00a313d8a05492f43346e7d9718b752672"


def check(value, why):
    if not value:
        raise AssertionError(why)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--attempt", default="r1")
    parser.add_argument("--with-ui", action="store_true")
    parser.add_argument("--with-office", action="store_true")
    args = parser.parse_args()
    path = ROOT / "fixtures/table-extraction/criteria.json"
    check(
        hashlib.sha256(path.read_bytes()).hexdigest() == CRITERIA_SHA, "fixed criteria"
    )
    fixed = json.loads(path.read_text(encoding="utf-8"))
    results = []
    cases = ["table-" + Path(case["file"]).stem + ".xlsx" for case in fixed["cases"]]
    if args.with_ui:
        cases.append("table-ui.xlsx")
    for name in cases:
        path = args.output / name
        row = {"file": name}
        try:
            book = openpyxl.load_workbook(path, data_only=False)
            check(len(book.worksheets) == 1, "one worksheet")
            sheet = book.worksheets[0]
            expected = [list(row) for row in fixed["expected_cells"]]
            if name == "table-ui.xlsx":
                expected[1][
                    1
                ] = "0099"  # Fixed UI keyboard input before its first execution.
            check(
                sheet.max_row == len(expected) and sheet.max_column == len(expected[0]),
                "dimensions",
            )
            for r, line in enumerate(expected, 1):
                for c, text in enumerate(line, 1):
                    cell = sheet.cell(r, c)
                    check((cell.value or "") == text, f"exact cell {r},{c}")
                    if text:
                        check(cell.data_type == "s", f"string cell {r},{c}")
                    check(cell.alignment.wrap_text is True, "multiline wrapping")
            with zipfile.ZipFile(path) as zipped:
                check(zipped.testzip() is None, "ZIP CRC")
                names = zipped.namelist()
                check(
                    not any(
                        "external" in n.lower() or "vba" in n.lower() for n in names
                    ),
                    "no external/macro parts",
                )
                xml = zipped.read("xl/worksheets/sheet1.xml").decode("utf-8")
                check("<f>" not in xml and "<f " not in xml, "no executable formulas")
            row.update(
                status="PASS", sha256=hashlib.sha256(path.read_bytes()).hexdigest()
            )
        except Exception as error:
            row.update(status="FAIL", error=str(error))
        results.append(row)
    if args.with_office:
        path = args.output / "table-calc.pdf"
        row = {"file": path.name}
        try:
            reader = PdfReader(path, strict=True)
            check(len(reader.pages) == 1, "Office PDF one page")
            first = reader.pages[0].extract_text()
            with pdfium.PdfDocument(path) as document:
                page = document[0]
                textpage = page.get_textpage()
                try:
                    second = textpage.get_text_range()
                    for line in fixed["expected_cells"]:
                        for text in line:
                            for phrase in text.split("\n"):
                                if phrase:
                                    check(
                                        phrase in first and phrase in second,
                                        "literal Office PDF text: " + phrase,
                                    )
                    check(
                        page.render(scale=1).to_pil().getbbox() is not None,
                        "Office PDFium render",
                    )
                finally:
                    textpage.close()
                    page.close()
            row.update(
                status="PASS", sha256=hashlib.sha256(path.read_bytes()).hexdigest()
            )
        except Exception as error:
            row.update(status="FAIL", error=str(error))
        results.append(row)
    report = {
        "criteria_sha256": CRITERIA_SHA,
        "cases": results,
        "failures": sum(r["status"] != "PASS" for r in results),
    }
    with (args.output / ("table-independent-" + args.attempt + ".json")).open(
        "x", encoding="utf-8"
    ) as file:
        json.dump(report, file, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return bool(report["failures"])


if __name__ == "__main__":
    raise SystemExit(main())
