"""Freeze synthetic DOCX input and expected content before conversion."""

import argparse
import hashlib
import json
from pathlib import Path
import zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image, ImageDraw


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    image = Image.new("RGB", (128, 64), "#3268b8")
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, 60, 50), fill="#e8b644")
    draw.rectangle((80, 10, 118, 50), fill="#4b8b5c")
    image.save(args.output / "fixture.png")
    texts = {
        "simple.docx": [
            [
                "日本語・English Document",
                "検索できる日本語PDFです。",
                "English document conversion preserves searchable text.",
                "項目",
                "金額",
                "合計",
                "1200円",
            ]
        ],
        "pages-images.docx": [
            ["1ページ目 First page", "画像と本文を保持します。"],
            ["2ページ目 Second page", "日本語とEnglishをコピーできます。"],
        ],
    }
    for name in texts:
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(20)
        section.left_margin = section.right_margin = Mm(20)
        style = doc.styles["Normal"]
        style.font.name = "Meiryo UI"
        style.font.size = Pt(11)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + kind), "Meiryo UI")
        doc.core_properties.author = "PDF達人の合成試験"
        doc.core_properties.title = "CC0 DOCX conversion fixture"
        if name == "simple.docx":
            for text in texts[name][0][:3]:
                doc.add_paragraph(text)
            table = doc.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            for cell, text in zip(
                [c for row in table.rows for c in row.cells], texts[name][0][3:]
            ):
                cell.text = text
        else:
            for text in texts[name][0]:
                doc.add_paragraph(text)
            doc.add_picture(str(args.output / "fixture.png"), width=Mm(40))
            doc.add_page_break()
            for text in texts[name][1]:
                doc.add_paragraph(text)
        doc.save(args.output / name)
    original = args.output / "simple.docx"
    for name, extra, data in (
        (
            "macro.docx",
            "word/vbaProject.bin",
            b"Refuse this synthetic macro marker before launch",
        ),
        (
            "embedded-object.docx",
            "word/embeddings/fixture.bin",
            b"Refuse synthetic embedded object",
        ),
        (
            "remote-image.docx",
            "word/_rels/document.xml.rels",
            b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdRemote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="http://127.0.0.1/never-fetch.png" TargetMode="External"/></Relationships>',
        ),
    ):
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(
            args.output / name, "x", zipfile.ZIP_DEFLATED
        ) as target:
            for entry in source.infolist():
                if entry.filename != extra:
                    target.writestr(entry, source.read(entry.filename))
            target.writestr(extra, data)
    (args.output / "invalid.docx").write_bytes(b"This is not a DOCX package.")
    files = {
        file.name: hashlib.sha256(file.read_bytes()).hexdigest()
        for file in sorted(args.output.iterdir())
    }
    fixed = dict(
        files=files,
        expected_pages=texts,
        page_size_pt=[595.3, 841.9],
        geometry_tolerance_pt=0.5,
        image_rgb_sha256=hashlib.sha256(image.tobytes()).hexdigest(),
        image_size=[128, 64],
        image_physical_width_pt=40 * 72 / 25.4,
        refused=[
            "macro.docx",
            "embedded-object.docx",
            "remote-image.docx",
            "invalid.docx",
        ],
        scope="Synthetic CC0 text, table and raster inputs; expectations fixed before first product conversion. No private document or embedded system font. Original contracts unchanged.",
    )
    (args.output / "criteria.json").write_text(
        json.dumps(fixed, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(fixed, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
