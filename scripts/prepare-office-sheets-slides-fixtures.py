"""Freeze synthetic XLSX/PPTX expectations before any product conversion."""

import argparse
import hashlib
import json
from pathlib import Path
import zipfile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Font
from openpyxl.worksheet.page import PageMargins
from pptx import Presentation
from pptx.util import Inches, Pt


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    source_image = (
        Path(__file__).resolve().parents[1] / "fixtures/office-import/fixture.png"
    )
    image = args.output / "fixture.png"
    image.write_bytes(source_image.read_bytes())
    book = Workbook()
    first = book.active
    first.title = "日本語"
    first.append(["日本語の表", "数量", "金額"])
    first.append(["項目一", 10, 1200])
    first.append(["項目二", 20, 2400])
    first.append(["合計", "=SUM(B2:B3)", "=SUM(C2:C3)"])
    second = book.create_sheet("English")
    second.append(["English worksheet", "Search and copy text"])
    second.append(["Image and print area"])
    graphic = SheetImage(str(image))
    graphic.width = 40 / 25.4 * 96
    graphic.height = graphic.width / 2
    second.add_image(graphic, "B4")
    for sheet in book:
        for row in sheet.iter_rows(min_row=1, max_row=8, min_col=1, max_col=3):
            for cell in row:
                cell.font = Font(name="Meiryo UI", size=12)
        for column in ("A", "B", "C"):
            sheet.column_dimensions[column].width = 22
        for row in range(1, 9):
            sheet.row_dimensions[row].height = 28
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.scale = 100
        sheet.page_margins = PageMargins(
            left=20 / 25.4,
            right=20 / 25.4,
            top=20 / 25.4,
            bottom=20 / 25.4,
            header=0,
            footer=0,
        )
        sheet.print_area = "A1:C8"
    book.save(args.output / "sheets.xlsx")
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(5.625)
    slide_text = [
        ["日本語のプレゼン資料", "検索とコピーを確認します。"],
        ["English presentation", "PowerPoint slides become searchable PDF."],
    ]
    for index, paragraphs in enumerate(slide_text):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(2))
        for at, value in enumerate(paragraphs):
            paragraph = (
                shape.text_frame.paragraphs[0]
                if at == 0
                else shape.text_frame.add_paragraph()
            )
            paragraph.text = value
            for run in paragraph.runs:
                run.font.name = "Meiryo UI"
                run.font.size = Pt(24 if at == 0 else 16)
        if index == 0:
            slide.shapes.add_picture(
                str(image), Inches(0.5), Inches(3), width=Inches(40 / 25.4)
            )
    presentation.save(args.output / "slides.pptx")
    for source, name, member in [
        ("sheets.xlsx", "macro.xlsx", "xl/vbaProject.bin"),
        ("slides.pptx", "embedded.pptx", "ppt/embeddings/fixture.bin"),
        ("sheets.xlsx", "external-book.xlsx", "xl/externalLinks/externalLink1.xml"),
    ]:
        with zipfile.ZipFile(args.output / source) as original, zipfile.ZipFile(
            args.output / name, "x", zipfile.ZIP_DEFLATED
        ) as changed:
            for entry in original.infolist():
                changed.writestr(entry, original.read(entry))
            changed.writestr(member, b"synthetic refused input")
    remote = Workbook()
    remote.active["A1"] = '=WEBSERVICE("http://127.0.0.1/never-fetch")'
    remote.save(args.output / "external-formula.xlsx")
    fixed = {
        "files": {
            p.name: digest(p) for p in sorted(args.output.iterdir()) if p.is_file()
        },
        "expected_pages": {
            "sheets.xlsx": [
                [
                    "日本語の表",
                    "数量",
                    "金額",
                    "項目一",
                    "項目二",
                    "合計",
                    "30",
                    "3600",
                ],
                ["English worksheet", "Search and copy text", "Image and print area"],
            ],
            "slides.pptx": slide_text,
        },
        "page_size_pt": {"sheets.xlsx": [595.3, 841.9], "slides.pptx": [720, 405]},
        "geometry_tolerance_pt": 0.5,
        "image_size": [128, 64],
        "image_rgb_sha256": "52fb5ac6cc0cd39e3b6845e337a15331f17465a639a08052a8371faf16e06d5b",
        "image_page": {"sheets.xlsx": 1, "slides.pptx": 0},
        "image_physical_width_pt": 40 * 72 / 25.4,
        "refused": [
            "macro.xlsx",
            "embedded.pptx",
            "external-book.xlsx",
            "external-formula.xlsx",
        ],
        "scope": "Synthetic CC0 text and graphic. Expectations fixed before product conversion; original contracts unchanged.",
    }
    (args.output / "criteria.json").write_text(
        json.dumps(fixed, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "files": len(fixed["files"]),
                "criteria_sha256": digest(args.output / "criteria.json"),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
